"""
High-Fidelity Legal Document (.docx) Generator for Affidavit in Reply.
Strictly conforms to High Court formatting conventions.
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from pathlib import Path
from typing import Dict, Any


def set_cell_margins(cell, top=50, bottom=50, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def create_affidavit_docx(data: Dict[str, Any], output_path: str) -> str:
    """
    Builds a beautifully typeset Affidavit in Reply .docx conforming strictly
    to the Bombay High Court formatting guidelines.
    """
    doc = Document()

    # Set page margins (Top 1", Bottom 1", Left 1.25", Right 1")
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.0)

    # Base Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.35
    normal_style.paragraph_format.space_after = Pt(6)

    case_details = data.get("case_details", {})
    deponent = data.get("deponent", {})
    body_paragraphs = data.get("body_paragraphs", [])
    prayers = data.get("prayers", [])
    attestation = data.get("attestation", {})
    advocate = data.get("advocate", {})

    def add_centered_heading(text: str, bold: bool = True, space_after: int = 4):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(2)
        run = p.add_run(text)
        run.bold = bold
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        return p

    # 1. Forum Heading (Bold, ALL CAPS, centered)
    add_centered_heading(case_details.get("court", "IN THE HIGH COURT OF JUDICATURE AT BOMBAY").upper(), bold=True)

    # 2. Jurisdiction (Bold, ALL CAPS, centered)
    jurisdiction = case_details.get("jurisdiction", "ORDINARY ORIGINAL CIVIL JURISDICTION").upper()
    if not jurisdiction.endswith("JURISDICTION"):
        jurisdiction += " JURISDICTION"
    add_centered_heading(jurisdiction, bold=True)

    # 3. Case Number (Bold, ALL CAPS, centered)
    proc_type = case_details.get("proceeding_type", "WRIT PETITION").upper()
    case_no = case_details.get("case_number", "1847")
    year = case_details.get("year", "2026")
    add_centered_heading(f"{proc_type} NO. {case_no} OF {year}", bold=True, space_after=14)

    # 4. Cause Title Table (Names left, status tags right, VERSUS centered)
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    def add_party_row(name_text: str, tag_text: str):
        row = table.add_row()
        cell_left, cell_right = row.cells[0], row.cells[1]
        cell_left.width = Inches(4.5)
        cell_right.width = Inches(1.75)
        set_cell_margins(cell_left, top=20, bottom=20, left=0, right=50)
        set_cell_margins(cell_right, top=20, bottom=20, left=50, right=0)
        
        p_left = cell_left.paragraphs[0]
        p_left.paragraph_format.space_after = Pt(2)
        r_l = p_left.add_run(name_text)
        r_l.font.name = 'Times New Roman'
        
        p_right = cell_right.paragraphs[0]
        p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_right.paragraph_format.space_after = Pt(2)
        r_r = p_right.add_run(tag_text)
        r_r.font.name = 'Times New Roman'

    petitioner_name = case_details.get("petitioner", "Sunrise Housing Private Limited")
    add_party_row(petitioner_name, "...Petitioner")

    # VERSUS on centered row
    row_v = table.add_row()
    cell_v = row_v.cells[0]
    cell_v.merge(row_v.cells[1])
    set_cell_margins(cell_v, top=40, bottom=40, left=0, right=0)
    p_v = cell_v.paragraphs[0]
    p_v.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_v.paragraph_format.space_after = Pt(4)
    p_v.paragraph_format.space_before = Pt(4)
    r_v = p_v.add_run("VERSUS")
    r_v.bold = True
    r_v.font.name = 'Times New Roman'

    # Respondents
    respondents = case_details.get("respondents", [])
    for idx, resp in enumerate(respondents, 1):
        add_party_row(f"{idx}. {resp}", f"...Respondent No.{idx}")

    # Spacer after cause title
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(10)

    # 5. Affidavit Title (Bold, ALL CAPS, centered)
    resp_num = case_details.get("answering_respondent_number", 2)
    add_centered_heading(f"AFFIDAVIT IN REPLY ON BEHALF OF RESPONDENT NO. {resp_num}", bold=True, space_after=12)

    # 6. Deponent Clause (One sentence, not numbered, justified)
    dep_name = deponent.get("name", "Arvind Rajan")
    designation = deponent.get("designation")
    organisation = deponent.get("organisation") or case_details.get("answering_respondent_name")
    address = deponent.get("address", "Bandra East, Mumbai, Maharashtra")
    verb = deponent.get("verification_verb", "solemnly affirm")

    p_dep = doc.add_paragraph()
    p_dep.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_dep.paragraph_format.space_after = Pt(10)
    
    if designation and organisation:
        dep_clause_text = (
            f"I, {dep_name}, having office at {address}, the {designation} of the "
            f"Respondent No.{resp_num} above named, do hereby {verb} and state as under:"
        )
    else:
        dep_clause_text = (
            f"I, {dep_name}, residing at {address}, the Respondent No.{resp_num} above named, "
            f"do hereby {verb} and state as under:"
        )
    r_dep = p_dep.add_run(dep_clause_text)
    r_dep.font.name = 'Times New Roman'

    # 7. Numbered Paragraphs
    for para_dict in body_paragraphs:
        p_num = para_dict.get("number")
        text = para_dict.get("text", "")
        p_body = doc.add_paragraph()
        p_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_body.paragraph_format.space_after = Pt(8)
        
        # Bold number prefix
        r_num = p_body.add_run(f"{p_num}. ")
        r_num.bold = True
        r_num.font.name = 'Times New Roman'

        # If text contains exhibit references, highlight/bold them
        if "EXHIBIT" in text:
            parts = text.split("EXHIBIT")
            p_body.add_run(parts[0]).font.name = 'Times New Roman'
            for part in parts[1:]:
                # Extract exhibit tag e.g. -'A'
                tokens = part.split(" ", 1)
                ex_run = p_body.add_run("EXHIBIT" + tokens[0])
                ex_run.bold = True
                ex_run.font.name = 'Times New Roman'
                if len(tokens) > 1:
                    p_body.add_run(" " + tokens[1]).font.name = 'Times New Roman'
        else:
            r_text = p_body.add_run(text)
            r_text.font.name = 'Times New Roman'

    # 8. Prayer (Bold caps heading, lettered items)
    add_centered_heading("PRAYER", bold=True, space_after=8)
    
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_after = Pt(6)
    r_intro = p_intro.add_run("I therefore respectfully pray that this Hon'ble Court may be pleased to:")
    r_intro.font.name = 'Times New Roman'

    for letter, prayer_item in zip(["(a)", "(b)", "(c)", "(d)"], prayers):
        p_pr = doc.add_paragraph()
        p_pr.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_pr.paragraph_format.left_indent = Inches(0.4)
        p_pr.paragraph_format.space_after = Pt(4)
        
        r_let = p_pr.add_run(f"{letter} ")
        r_let.bold = True
        r_let.font.name = 'Times New Roman'
        
        r_pit = p_pr.add_run(f"{prayer_item}")
        r_pit.font.name = 'Times New Roman'

    # 9. Jurat (Solemnly affirmed, DEPONENT right, Before Me left)
    p_jurat_space = doc.add_paragraph()
    p_jurat_space.paragraph_format.space_after = Pt(10)

    jurat_verb_cap = "Solemnly affirmed" if "solemnly" in verb.lower() else "Sworn"
    place = attestation.get("place", "Mumbai")
    date_ord = attestation.get("date_ordinal", "5th day of September 2026")

    j_table = doc.add_table(rows=2, cols=2)
    j_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    j_table.autofit = False
    
    j_table.rows[0].cells[0].width = Inches(3.5)
    j_table.rows[0].cells[1].width = Inches(2.75)
    j_table.rows[1].cells[0].width = Inches(3.5)
    j_table.rows[1].cells[1].width = Inches(2.75)

    cell_j1 = j_table.rows[0].cells[0].paragraphs[0]
    cell_j1.paragraph_format.space_after = Pt(2)
    r1 = cell_j1.add_run(f"{jurat_verb_cap} at {place}\nOn this {date_ord}")
    r1.font.name = 'Times New Roman'

    cell_j2 = j_table.rows[0].cells[1].paragraphs[0]
    cell_j2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    cell_j2.paragraph_format.space_after = Pt(2)
    r2 = cell_j2.add_run("DEPONENT")
    r2.bold = True
    r2.font.name = 'Times New Roman'

    cell_j3 = j_table.rows[1].cells[0].paragraphs[0]
    cell_j3.paragraph_format.space_before = Pt(12)
    r3 = cell_j3.add_run("Before Me")
    r3.font.name = 'Times New Roman'

    # 10. Verification (Heading bold caps centered, range matching body paragraphs)
    add_centered_heading("VERIFICATION", bold=True, space_after=8)

    total_body_paras = len(body_paragraphs)
    p_ver = doc.add_paragraph()
    p_ver.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_ver.paragraph_format.space_after = Pt(8)
    
    ver_text = (
        f"I, {dep_name}, the Deponent above named, do hereby verify that the contents of "
        f"paragraphs 1 to {total_body_paras} and the Prayer above are true and correct to "
        f"my knowledge and belief and that nothing material has been concealed therefrom."
    )
    r_ver = p_ver.add_run(ver_text)
    r_ver.font.name = 'Times New Roman'

    # Verification date and deponent block
    v_table = doc.add_table(rows=1, cols=2)
    v_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    v_table.rows[0].cells[0].width = Inches(4.0)
    v_table.rows[0].cells[1].width = Inches(2.25)
    
    p_v1 = v_table.rows[0].cells[0].paragraphs[0]
    r_v1 = p_v1.add_run(f"Verified at {place} on this {date_ord}.")
    r_v1.font.name = 'Times New Roman'

    p_v2 = v_table.rows[0].cells[1].paragraphs[0]
    p_v2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_v2 = p_v2.add_run("DEPONENT")
    r_v2.bold = True
    r_v2.font.name = 'Times New Roman'

    # Advocate / Drafting Block
    adv_firm = advocate.get("firm_name", "Rajan & Associates")
    adv_for = advocate.get("acting_for", f"Respondent No.{resp_num}")
    
    p_adv_space = doc.add_paragraph()
    p_adv_space.paragraph_format.space_after = Pt(14)
    
    p_adv = doc.add_paragraph()
    p_adv.paragraph_format.space_after = Pt(2)
    r_adv1 = p_adv.add_run(adv_firm.upper())
    r_adv1.bold = True
    r_adv1.font.name = 'Times New Roman'

    p_adv2 = doc.add_paragraph()
    r_adv2 = p_adv2.add_run(f"Advocates for the {adv_for}.")
    r_adv2.font.name = 'Times New Roman'

    # Save document
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path
